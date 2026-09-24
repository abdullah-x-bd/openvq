#!/usr/bin/env python3
"""Prepare condition-stratified TCD-VoIP manifests for OpenVQ validation."""
import argparse
import csv
import json
import re
from collections import defaultdict
from pathlib import Path
import openpyxl

def family_from_name(name):
    m = re.match(r"^[CR]_\d+_([^_]+)_", name)
    return m.group(1).upper() if m else "UNKNOWN"

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("root")
    ap.add_argument("outdir")
    args=ap.parse_args()
    root=Path(args.root)
    out=Path(args.outdir)
    out.mkdir(parents=True, exist_ok=True)

    xlsx=next(root.rglob("*.xlsx"))
    wb=openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws=wb["Subjective Test Scores"]

    wav_by_name={}
    for p in root.rglob("*.wav"):
        wav_by_name[p.name]=p.resolve()

    rows=[]
    headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
    idx={h:i for i,h in enumerate(headers)}
    for values in ws.iter_rows(min_row=2, values_only=True):
        filename=values[idx["Filename"]]
        if not filename:
            continue
        condition=int(values[idx["ConditionID"]])
        mos=float(values[idx["sample MOS"]])
        degraded=wav_by_name.get(filename)
        reference=wav_by_name.get("R_"+filename[2:]) if filename.startswith("C_") else None
        if degraded is None or reference is None:
            raise SystemExit(f"missing pair for {filename}: {degraded=} {reference=}")
        listener_scores=[
            float(values[i]) for i,h in enumerate(headers)
            if isinstance(h,int) and values[i] is not None
        ]
        rows.append({
            "reference":str(reference),
            "degraded":str(degraded),
            "human_mos":mos,
            "condition_id":condition,
            "family":family_from_name(filename),
            "filename":filename,
            "listener_n":len(listener_scores),
            "listener_sd":(
                (sum((x-(sum(listener_scores)/len(listener_scores)))**2 for x in listener_scores)
                 / max(1,len(listener_scores)-1))**0.5
                if listener_scores else 0.0
            ),
        })

    # Split by condition ID within degradation family, never by individual clip.
    by_family=defaultdict(list)
    for r in rows:
        if r["condition_id"] not in by_family[r["family"]]:
            by_family[r["family"]].append(r["condition_id"])
    split_of={}
    for family, ids in sorted(by_family.items()):
        ids=sorted(ids)
        n=len(ids)
        # Deterministic approximately 70/15/15, with at least one dev/test if possible.
        n_test=max(1, round(n*0.15)) if n>=3 else 0
        n_dev=max(1, round(n*0.15)) if n>=4 else 0
        test=set(ids[-n_test:]) if n_test else set()
        dev=set(ids[-(n_test+n_dev):-n_test] if n_test else ids[-n_dev:]) if n_dev else set()
        for cid in ids:
            split_of[(family,cid)]="test" if cid in test else ("dev" if cid in dev else "train")

    fields=["reference","degraded","human_mos","condition_id","family","filename","listener_n","listener_sd"]
    for split in ["all","train","dev","test"]:
        chosen=rows if split=="all" else [r for r in rows if split_of[(r["family"],r["condition_id"])]==split]
        with open(out/f"{split}.csv","w",newline="",encoding="utf-8") as f:
            w=csv.DictWriter(f,fieldnames=fields)
            w.writeheader()
            w.writerows(chosen)

    summary={
        "total_clips":len(rows),
        "families":{},
        "splits":{},
        "pairing":"C_... degraded paired with R_... reference by filename",
    }
    for family in sorted(by_family):
        fr=[r for r in rows if r["family"]==family]
        summary["families"][family]={
            "clips":len(fr),
            "conditions":len(set(r["condition_id"] for r in fr)),
            "mean_mos":sum(r["human_mos"] for r in fr)/len(fr),
        }
    for split in ["train","dev","test"]:
        rr=[r for r in rows if split_of[(r["family"],r["condition_id"])]==split]
        summary["splits"][split]={
            "clips":len(rr),
            "conditions":len(set((r["family"],r["condition_id"]) for r in rr)),
            "families":dict(sorted(__import__("collections").Counter(r["family"] for r in rr).items()))
        }
    (out/"split-summary.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")
    print(json.dumps(summary,indent=2))

if __name__=="__main__":
    main()
