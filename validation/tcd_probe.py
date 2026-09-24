#!/usr/bin/env python3
import json
import sys
import wave
from collections import Counter, defaultdict
from pathlib import Path
from zipfile import ZipFile

root = Path(sys.argv[1])
wavs = sorted(root.rglob("*.wav"))
xlsx = sorted(root.rglob("*.xlsx"))
txt = sorted(root.rglob("*.txt"))

result = {
    "wav_count": len(wavs),
    "xlsx": [str(p.relative_to(root)) for p in xlsx],
    "txt": [str(p.relative_to(root)) for p in txt],
    "first_80_wavs": [str(p.relative_to(root)) for p in wavs[:80]],
    "extensions": dict(Counter(p.suffix.lower() for p in root.rglob("*") if p.is_file())),
}

# Inspect audio format distribution.
formats = Counter()
durations = []
for p in wavs:
    try:
        with wave.open(str(p), "rb") as w:
            formats[(w.getframerate(), w.getnchannels(), w.getsampwidth())] += 1
            durations.append(w.getnframes() / w.getframerate())
    except Exception as e:
        formats[("error", type(e).__name__, 0)] += 1

result["wav_formats"] = {str(k): v for k, v in formats.items()}
if durations:
    result["duration_s"] = {
        "min": min(durations),
        "max": max(durations),
        "mean": sum(durations) / len(durations),
    }

if xlsx:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx[0], read_only=True, data_only=True)
    result["sheets"] = wb.sheetnames
    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(min_row=1, max_row=min(12, ws.max_row), values_only=True))
        sheets[name] = {
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "sample_rows": rows,
        }
    result["sheet_probe"] = sheets

# Heuristics for reference-like file names.
ref_words = ("ref", "clean", "source", "original", "undist", "anchor")
result["reference_like"] = [
    str(p.relative_to(root)) for p in wavs
    if any(w in p.name.lower() for w in ref_words)
][:100]

# Group names by first underscore-delimited token to understand naming.
groups = defaultdict(list)
for p in wavs:
    groups[p.stem.split("_")[0]].append(p.name)
result["prefix_groups"] = {
    k: v[:12] for k, v in sorted(groups.items(), key=lambda kv: (-len(kv[1]), kv[0]))[:30]
}

print(json.dumps(result, indent=2, default=str))
